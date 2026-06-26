"""Output writers that persist Whisper segments to TXT, SRT, and DOCX formats."""

from __future__ import annotations

from datetime import timedelta
from pathlib import Path
from typing import Any


def format_timestamp(seconds: float) -> str:
    """Convert a duration in seconds to an SRT timestamp (HH:MM:SS,mmm).

    Args:
        seconds: A non-negative duration in seconds.

    Returns:
        The formatted timestamp string, e.g. "00:01:23,456".
    """
    total_ms = int(timedelta(seconds=seconds).total_seconds() * 1000)
    hours, remainder = divmod(total_ms, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    secs, millis = divmod(remainder, 1_000)
    return f"{hours:02d}:{minutes:02d}:{secs:02d},{millis:03d}"


def write_srt(segments: list[dict[str, Any]], srt_path: Path) -> None:
    """Write Whisper segments to a SubRip (.srt) subtitle file.

    Args:
        segments: Whisper transcription segments with 'start', 'end', and 'text'.
        srt_path: Destination path for the subtitle file.
    """
    srt_path.parent.mkdir(parents=True, exist_ok=True)
    with srt_path.open("w", encoding="utf-8") as handle:
        for index, segment in enumerate(segments, start=1):
            start = format_timestamp(segment["start"])
            end = format_timestamp(segment["end"])
            text = segment["text"].strip()
            handle.write(f"{index}\n{start} --> {end}\n{text}\n\n")


def write_txt(segments: list[dict[str, Any]], txt_path: Path) -> None:
    """Write the plain-text transcript, one segment per line.

    Args:
        segments: Whisper transcription segments.
        txt_path: Destination path for the text file.
    """
    txt_path.parent.mkdir(parents=True, exist_ok=True)
    with txt_path.open("w", encoding="utf-8") as handle:
        for segment in segments:
            handle.write(segment["text"].strip() + "\n")


def write_docx(
    segments: list[dict[str, Any]],
    docx_path: Path,
    include_timestamps: bool = True,
) -> None:
    """Export the transcript to a formatted Word (.docx) document.

    Args:
        segments:           Whisper transcription segments.
        docx_path:          Destination path for the document.
        include_timestamps: Prefix each paragraph with its time range when True.

    Note:
        Requires the optional `python-docx` dependency.
    """
    from docx import Document
    from docx.shared import Pt

    docx_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(11)

    document.add_heading("Transcript", level=1)
    for segment in segments:
        text = segment["text"].strip()
        prefix = ""
        if include_timestamps:
            prefix = (
                f"[{format_timestamp(segment['start'])} – "
                f"{format_timestamp(segment['end'])}] "
            )
        document.add_paragraph(prefix + text)

    document.save(str(docx_path))